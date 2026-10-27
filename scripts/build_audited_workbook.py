"""Build a fully reproducible audit workbook using the public openpyxl package."""
from __future__ import annotations

import csv
import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
OUT = ROOT / "outputs"

def read(path):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))

regions = read(DATA / "regions_2025.csv")
sources = read(DATA / "sources.csv")
snapshots = read(DATA / "source_snapshot_manifest.csv")
facts = read(DATA / "observations_reviewed.csv")
audits = read(OUT / "source_number_audit.csv")
analysis = json.loads((OUT / "analysis_results.json").read_text(encoding="utf-8"))
wb = Workbook()
wb.remove(wb.active)

def sheet(name, headers, rows):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 28
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="16324F")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True)
    for idx, header in enumerate(headers, 1):
        width = min(45, max(14, len(str(header)) * 2 + 4))
        ws.column_dimensions[get_column_letter(idx)].width = width
    return ws

sheet("使用说明", ["事项", "说明"], [
    ["研究状态", "2025养老金三类金额齐全2县，人数也齐全1县；两县养老金来源均为转载，不能外推全省。"],
    ["快照审计", "成功归档15/16个登记URL；数值定位不等于口径核实。详见来源快照与逐条审计。"],
    ["公式", "养老金月均代理=全年发放万元×10000÷年末领取人数÷12；非实际人均月待遇。"],
    ["复现", "python scripts/prepare_data.py; python scripts/reproduce.py; python scripts/audit_snapshots.py; python scripts/build_audited_workbook.py"],
    ["旧版数据", "上一版基础数据库.xlsx保留作线索，未逐条核验的数据不进入本工作簿事实层。"],
    ["资金边界", "社保基金、一般公共预算、住房公积金分别核算；禁止跨账相减。"],
])
sheet("行政区100", list(regions[0]), [[r[k] for k in regions[0]] for r in regions])
sheet("来源索引", list(sources[0]), [[r[k] for k in sources[0]] for r in sources])
sheet("来源快照", list(snapshots[0]), [[r[k] for k in snapshots[0]] for r in snapshots])
fact_headers = list(facts[0]) + ["source_version"]
sheet("事实142条", fact_headers, [
    [float(r[k]) if k == "value" else r[k] for k in facts[0]] + [r["source_id"]]
    for r in facts
])
sheet("逐条来源审计", list(audits[0]), [[r[k] for k in audits[0]] for r in audits])
indicators = sorted({r["indicator_id"] for r in facts})
definition = {
    "resident_population_year_end": "年末常住人口；县表原始单位万人，转为人",
    "urbanization_rate": "年末常住人口城镇化率，百分数",
    "hukou_age_60plus": "年末户籍60岁及以上人口，不可除以常住人口",
    "births": "当年出生人口；自然变动=出生-死亡",
    "deaths": "当年死亡人口",
    "provident_contributors": "住房公积金年度实缴人数，非县级就业总人数",
    "urban_nonprivate_avg_wage": "城镇非私营单位就业人员平均工资",
    "urban_private_avg_wage": "城镇私营单位就业人员平均工资",
}
for group in ("enterprise", "public", "resident"):
    definition[f"pension_{group}_recipients_year_end"] = "年末待遇领取人数；不是全年平均人数"
    definition[f"pension_{group}_annual_paid"] = "全年发放养老金金额；公报口径待与经办机构核对"
    definition[f"pension_{group}_active_or_paid"] = "在职参保或居民实际缴费人数；两者不同口径"
    definition[f"pension_{group}_contribution"] = "征缴基金，不等于待遇全部资金来源"
sheet("指标口径", ["indicator_id", "定义与限制", "是否完整终核"], [
    [i, definition.get(i, "详见来源公报及事实备注；尚未形成统一统计定义"), "否" if i not in definition else "按来源限定"]
    for i in indicators
])

ws = sheet("峡江养老金复算", ["年份", "制度", "年末领取人数", "年度发放万元", "人数占比", "金额占比", "月均金额代理元", "来源", "证据等级"], [
    [year, group["name"], group["recipients_year_end"], group["annual_paid_wanyuan"], None, None, None, "S03" if year == 2024 else "S02", "B_转载待官方原件"]
    for year, result in [(2024, analysis["xiajiang_2024"]), (2025, analysis["xiajiang_2025"])]
    for group in result["groups"]
])
for i in range(2, 8):
    start = 2 if i <= 4 else 5
    end = start + 2
    ws[f"E{i}"] = f"=C{i}/SUM(C${start}:C${end})"
    ws[f"F{i}"] = f"=D{i}/SUM(D${start}:D${end})"
    ws[f"G{i}"] = f"=D{i}*10000/C{i}/12"
    for col in "EF":
        ws[f"{col}{i}"].number_format = "0.00%"

ws = sheet("养老金金额分解", ["制度", "2024金额万元", "2025金额万元", "变化万元", "人数分量万元", "金额人均代理分量万元", "恒等误差", "解释"], [
    [name, None, None, None, None, None, None, "对称分解；年末人数代理；非因果效应"]
    for name in ["企业职工", "机关事业", "城乡居民"]
])
for i in range(2, 5):
    a, b = i, i + 3
    ws[f"B{i}"] = f"='峡江养老金复算'!D{a}"
    ws[f"C{i}"] = f"='峡江养老金复算'!D{b}"
    ws[f"D{i}"] = f"=C{i}-B{i}"
    ws[f"E{i}"] = f"=('峡江养老金复算'!C{b}-'峡江养老金复算'!C{a})*(B{i}/'峡江养老金复算'!C{a}+C{i}/'峡江养老金复算'!C{b})/2"
    ws[f"F{i}"] = f"=(C{i}/'峡江养老金复算'!C{b}-B{i}/'峡江养老金复算'!C{a})*('峡江养老金复算'!C{a}+'峡江养老金复算'!C{b})/2"
    ws[f"G{i}"] = f"=D{i}-E{i}-F{i}"

ws = sheet("赣州18县人口", ["行政代码", "县区", "2024常住人口人", "2025常住人口人", "人口变化人", "2024城镇化率%", "2025城镇化率%", "城镇化变化百分点", "证据"], [
    [r["region_code"], r["county"], r["pop_2024"], r["pop_2025"], None, r["urbanization_2024"], r["urbanization_2025"], None, "S06/S05政府表格单元格复核"]
    for r in analysis["ganzhou_18_counties"]
])
for i in range(2, 20):
    ws[f"E{i}"] = f"=D{i}-C{i}"
    ws[f"H{i}"] = f"=G{i}-F{i}"

by_indicator = Counter((r["year"], r["indicator_id"]) for r in facts if r["level"] == "县")
key_indicators = ["resident_population_year_end", "urbanization_rate", "pension_enterprise_annual_paid", "pension_public_annual_paid", "pension_resident_annual_paid", "pension_enterprise_recipients_year_end", "pension_public_recipients_year_end", "pension_resident_recipients_year_end", "provident_contributors"]
sheet("县级覆盖与缺失", ["年份", "指标", "有值县数", "缺失县数", "缺失率", "主要原因"], [
    [year, indicator, by_indicator[(str(year), indicator)], 100 - by_indicator[(str(year), indicator)], (100 - by_indicator[(str(year), indicator)]) / 100, "公开来源未采得可核县级同口径值；空值不表示0"]
    for year in (2024, 2025) for indicator in key_indicators
])
for row in wb["县级覆盖与缺失"].iter_rows(min_row=2):
    row[4].number_format = "0%"

target = OUT / "江西县域社会保障_2025可复现审计工作簿.xlsx"
wb.save(target)
print(target)
